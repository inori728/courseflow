"""把'表格'变成课表数据(导入通道的解析核心).

迭代 1 覆盖最常用的一类输入:
- 从教务系统网页 / Excel 复制出来的'格子文字'(粘贴入口);
- .xlsx / .csv 文件(文件入口).

两者共用同一套'网格解析':先认表头(周一...周日),再从每个格子里
拆出'课名 / 教室 / 周次'.格式千奇百怪的情况,留给后面的'导入校对'版本.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Optional

from .model import Course, parse_day, parse_periods

# 周次文字:1-16周 / 第2周 / 4-11周 / 单周 / 双周 / 每周 ...
WEEK_RE = re.compile(
    r"(第?\s*[0-9一二三四五六七八九十]+\s*[-\u2013\u2014~至]\s*[0-9一二三四五六七八九十]+\s*周)"
    r"|(第?\s*[0-9一二三四五六七八九十]+\s*周)"
    r"|(单周)|(双周)|(每周)|(全周)"
)

# 拆分"粘连词"时要剥掉的边缘字符(括号,标点等)
_EDGE = "()（）[]【】{}<>《》·,，、;；:：… "


def _is_week_text(t: str) -> bool:
    return bool(WEEK_RE.search(t))


def _looks_like_room(t: str) -> bool:
    """猜一个词是不是教室.宁缺勿错--猜错会把课名弄乱."""
    if not t or len(t) > 12:
        return False
    has_digit = bool(re.search(r"\d", t))
    if has_digit and re.search(r"[楼馆室厅栋区场教]", t):
        return True                       # 信息楼302 / 机房A区101 / 教一A101
    if has_digit and re.search(r"[-\u2014#]", t):
        return True                       # 东4-201 / 11-303
    if re.fullmatch(r"[东南西北中]?[A-Za-z]?\d{3,5}", t):
        return True                       # 西4104 / 4104 / B203
    if re.search(r"(操场|体育馆|游泳馆|礼堂|报告厅|活动中心|会议室)$", t):
        return True                       # 风雨操场 / 校游泳馆
    return False


def _tokenize_cell(cell: str) -> list[str]:
    """把格子文字切成小词;顺手把"粘连的周次"拆开,如 高等数学(1-16周)."""
    text = re.sub(r"[\uff1b;]", " ", str(cell).replace("\r", ""))
    tokens: list[str] = []
    for t in text.split():
        if t.strip(_EDGE) == "":
            continue
        m = WEEK_RE.search(t)
        if m and m.group() != t:
            head = t[: m.start()].strip().strip(_EDGE)
            tail = t[m.end():].strip().strip(_EDGE)
            if head:
                tokens.append(head)
            tokens.append(m.group())
            if tail:
                tokens.append(tail)
        else:
            tokens.append(t)
    return tokens


def _parse_cell(cell: str) -> list[tuple[str, str, str]]:
    """解析一个格子 → [(课名, 教室, 周次), ...].

    规则:连续的小词里,出现"周次"之后再冒出新的"非周次非教室"词,
    就认为开始写下一门课了(很多教务系统同一格塞两门课).
    """
    tokens = _tokenize_cell(cell)
    if not tokens:
        return []

    segments: list[list[str]] = [[]]
    for t in tokens:
        seg = segments[-1]
        if seg and not _is_week_text(t) and not _looks_like_room(t) and any(
            _is_week_text(x) for x in seg
        ):
            segments.append([])
        segments[-1].append(t)

    result: list[tuple[str, str, str]] = []
    for seg in segments:
        name_parts: list[str] = []
        room = ""
        weeks = ""
        for t in seg:
            if _is_week_text(t):
                weeks = weeks or t
            elif _looks_like_room(t):
                room = room or t
            else:
                name_parts.append(t)
        if name_parts:
            result.append((" ".join(name_parts).strip(), room, weeks))
    return result


def parse_grid(rows: list[list]) -> list[Course]:
    """rows = 二维表格(第一行一般是表头).认不出表头时抛 ValueError."""
    rows = [["" if c is None else str(c).strip() for c in r] for r in rows]

    header_i: Optional[int] = None
    day_cols: dict[int, int] = {}
    for i, r in enumerate(rows[:10]):
        cols: dict[int, int] = {}
        for j, cell in enumerate(r):
            if not cell or len(cell) > 8:      # 表头格一般就"周一"这么短
                continue
            d = parse_day(cell)
            if d is not None and d not in cols:
                cols[d] = j
        if len(cols) >= 3:
            header_i, day_cols = i, cols
            break

    if header_i is None:
        raise ValueError("没认出来课表格子--需要有一行表头写着'周一...周日'")

    first_day_col = min(day_cols.values())
    if first_day_col < 1:
        raise ValueError("表头左边应该还有一列'节次',没找到")
    label_col = first_day_col - 1

    courses: list[Course] = []
    for r in rows[header_i + 1:]:
        if not any(c for c in r):
            continue
        periods = parse_periods(r[label_col] if len(r) > label_col else "")
        if periods is None:
            continue
        start, end = periods
        for day, col in day_cols.items():
            if col >= len(r) or not r[col]:
                continue
            for name, room, weeks in _parse_cell(r[col]):
                courses.append(
                    Course(name=name, day=day, start=start,
                           end=min(end, 12), room=room, weeks=weeks)
                )
    if not courses:
        raise ValueError("表头认出来了,但格子里没找到课程内容")
    return courses


def _tsv_rows(text: str) -> list[list[str]]:
    """把'表格文字'切成行列.兼容 Excel 复制出来的带引号换行格式."""
    rows: list[list[str]] = []
    row: list[str] = []
    cell: list[str] = []
    in_quote = False
    i = 0
    n = len(text)
    while i < n:
        ch = text[i]
        if in_quote:
            if ch == '"':
                if i + 1 < n and text[i + 1] == '"':
                    cell.append('"')
                    i += 2
                    continue
                in_quote = False
                i += 1
                continue
            cell.append(ch)
            i += 1
            continue
        if ch == '"':
            in_quote = True
            i += 1
            continue
        if ch == "\t":
            row.append("".join(cell))
            cell = []
        elif ch == "\n":
            row.append("".join(cell))
            cell = []
            rows.append(row)
            row = []
        elif ch == "\r":
            pass
        else:
            cell.append(ch)
        i += 1
    row.append("".join(cell))
    rows.append(row)
    return rows


def parse_pasted_text(text: str) -> list[Course]:
    """解析剪贴板里的'格子文字'."""
    if not text or not text.strip():
        raise ValueError("剪贴板里没有文字.先选中课表,Ctrl+C 复制,再来运行这个命令")

    rows = _tsv_rows(text)
    if all(len(r) <= 1 for r in rows):
        # 没看到制表符:按'连续空格'再切一次;
        # 表头那种 "周一 周二 周三 ..."(单空格)也顺手切开
        new_rows: list[list[str]] = []
        for r in rows:
            line = (r[0] if r else "").strip()
            parts = [p for p in re.split(r"\s{2,}", line) if p] if line else []
            if len(parts) <= 1 and line:
                days = sum(1 for p in line.split() if parse_day(p) is not None)
                if days >= 3:
                    parts = line.split()
            new_rows.append(parts)
        rows = new_rows

    return parse_grid(rows)


def parse_table_file(path) -> list[Course]:
    """从 .xlsx / .csv 文件读课表."""
    p = Path(path)
    if not p.exists():
        raise ValueError(f"没找到文件:{p}")
    suffix = p.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(p)
    if suffix == ".csv":
        with p.open(encoding="utf-8-sig", newline="") as f:
            rows = [r for r in csv.reader(f)]
        return parse_grid(rows)
    raise ValueError(f"暂不支持这种文件:{suffix}(现在支持 .xlsx / .csv;"
                     f"Excel 也可以先选中表格'复制→粘贴')")


def _parse_xlsx(p: Path) -> list[Course]:
    from openpyxl import load_workbook

    wb = load_workbook(p, data_only=True)
    best: Optional[list[Course]] = None
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        try:
            got = parse_grid(rows)
        except ValueError:
            continue
        if best is None or len(got) > len(best):
            best = got
    if best is None:
        raise ValueError("这个 Excel 里没找到课表格子(需要有一行表头写着'周一...周日')")
    return best


def read_clipboard_text() -> str:
    """读剪贴板里的文字(借 Python 自带的 tkinter,不装额外库)."""
    try:
        import tkinter
    except ImportError:
        return ""
    try:
        root = tkinter.Tk()
        root.withdraw()
        try:
            return str(root.clipboard_get())
        finally:
            root.destroy()
    except Exception:
        return ""
